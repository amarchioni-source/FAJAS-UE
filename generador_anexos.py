#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Anexos V (Declaracion Jurada de Fajas UE - SENASA)
================================================================

Toma una o varias "bajadas" (export de scans, hoja 'ag-grid') y genera UN anexo
por cada fecha de produccion, con el formato exacto de la plantilla, mas un PDF
cronologico (una pagina por dia) y un ZIP de los Excel individuales organizados
en carpetas por mes.

USO
---
    python generador_anexos.py <bajada1.xlsx> [<bajada2.xlsx> ...] \
        --plantilla plantilla_Fajas_UE.xlsx \
        --salida ./salida

  - Acepta archivos sueltos o una carpeta (toma todos los .xlsx que haya dentro).
  - Si no se pasa --plantilla, busca 'plantilla_Fajas_UE.xlsx' en el directorio del script.

REGLAS (resumen; detalle completo en LEEME.md)
----------------------------------------------
  * Mercado UE  = fajas cuyo serial empieza con F y la 2da letra es E o F (Serie E / Serie F).
                  Las FC (Serie C = USA/Kosher) se EXCLUYEN.
  * Serie       = 2da letra del serial (FE -> E, FF -> F).
  * Un anexo por fecha de produccion; todas las series del dia en el mismo cuadro.
  * Tramos      = dentro de cada serie, se separa en un tramo nuevo cuando hay un
                  salto > 1000 entre fajas consecutivas.
  * Solicitadas : Desde = primera faja usada del tramo ; Hasta = redondeo al millar
                  superior ; Cantidad = Hasta - Desde (formula).
  * Utilizadas  : Desde/Hasta = primera/ultima faja usada del tramo ; la Cantidad
                  (una sola, en la 1ra fila) = total de fajas FE/FF distintas del dia.
  * Devueltas   : Desde = Utilizada-Hasta ; Hasta = Solicitada-Hasta (cola sin usar).
  * Rotas       : una fila, Cantidad por formula de balance = SUM(Solicitadas) -
                  Utilizadas - SUM(Devueltas).
  * Total       : Cantidad = SUM(Solicitadas) (formula).
  * Pie         : "SERIE ___ Nº: ___" en blanco, se completa A MANO al firmar
                  (los numeros de precinto no estan en la bajada).

SALIDA
------
    salida/
      xlsx/                              <- un .xlsx por dia (nombre DD-MM-AAAA)
      Anexos_V_Fajas_UE_cronologico.pdf  <- todos los dias, una pagina c/u, por fecha
      Anexos_V_Fajas_UE_editables.zip    <- xlsx en carpetas 02_Febrero, 03_Marzo, ...

DEPENDENCIAS:  pandas, openpyxl, pypdf, y LibreOffice (soffice) para el PDF.
"""

import argparse, copy, glob, os, re, shutil, subprocess, sys
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.worksheet.properties import PageSetupProperties

GAP = 1000                       # salto que abre un tramo nuevo
SERIES_UE = {'E', 'F'}           # 2da letra del serial que cuenta como UE (FC excluida)
MESES = {'01':'01_Enero','02':'02_Febrero','03':'03_Marzo','04':'04_Abril',
         '05':'05_Mayo','06':'06_Junio','07':'07_Julio','08':'08_Agosto',
         '09':'09_Septiembre','10':'10_Octubre','11':'11_Noviembre','12':'12_Diciembre'}

# --------------------------------------------------------------------------- #
def parse_faja(x):
    """Devuelve (serie, numero) si el serial es F + letra-UE + digitos; si no, None."""
    x = str(x)
    m = re.match(r'^F([A-Z])(\d+)$', x)
    if m and m.group(1) in SERIES_UE:
        return m.group(1), int(m.group(2))
    return None

def day_tramos(d):
    """Para un DataFrame de un dia: lista de tramos + total de fajas UE usadas."""
    series = {}
    for col in ('Faja1', 'Faja2'):
        for v in d[col].dropna():
            p = parse_faja(v)
            if p:
                series.setdefault(p[0], set()).add(p[1])
    tramos, total = [], 0
    for letter in sorted(series):
        arr = np.array(sorted(series[letter]))
        total += len(arr)
        splits = np.where(np.diff(arr) > GAP)[0]
        start, segs = 0, []
        for s in splits:
            segs.append(arr[start:s+1]); start = s+1
        segs.append(arr[start:])
        for seg in segs:
            hu = int(seg[-1])
            tramos.append(dict(serie=letter, du=int(seg[0]), hu=hu,
                               sh=((hu // 1000) + 1) * 1000))   # redondeo al millar superior
    return tramos, total

# --------------------------------------------------------------------------- #
def _style(c):
    return dict(font=copy.copy(c.font), fill=copy.copy(c.fill), border=copy.copy(c.border),
                alignment=copy.copy(c.alignment), number_format=c.number_format)
def _apply(c, s, val=None):
    c.font=copy.copy(s['font']); c.fill=copy.copy(s['fill']); c.border=copy.copy(s['border'])
    c.alignment=copy.copy(s['alignment']); c.number_format=s['number_format']
    if val is not None: c.value = val

def build_day(plantilla, fecha, tramos, util_total, outpath):
    yy, mm, dd = fecha.split('-'); dia = int(dd)
    wb = openpyxl.load_workbook(plantilla); ws = wb['FAJAS']
    st = {k: [_style(ws[f'{c}{r}']) for c in 'ABCDE'] for k, r in
          {'hdr':15,'sol':16,'uti':18,'rot':20,'dev':21,'tot':23}.items()}
    foot_s = _style(ws['A28'])
    sig = [_style(ws[c]) for c in ('A36','A37','H35','H36','H37')]
    ws['F14'] = dia; ws['H14'] = f'Mes {mm}'; ws['J14'] = int(yy)
    for m in [mr for mr in list(ws.merged_cells.ranges) if mr.min_row >= 15]:
        ws.unmerge_cells(str(m))
    ws.delete_rows(15, ws.max_row - 14 + 5)
    r = 15
    for i, c in enumerate('ABCDE'):
        _apply(ws[f'{c}{r}'], st['hdr'][i], ['','Serie','Cantidad','Desde nº','Hasta nº'][i])
    ws.row_dimensions[r].height = 25.25
    sol_rows, dev_rows = [], []
    first = True
    for t in tramos:
        r += 1; sol_rows.append(r); ws.row_dimensions[r].height = 25.25
        _apply(ws[f'A{r}'], st['sol'][0], 'Solicitadas' if first else None); first = False
        _apply(ws[f'B{r}'], st['sol'][1], t['serie']); _apply(ws[f'C{r}'], st['sol'][2], f'=+E{r}-D{r}')
        _apply(ws[f'D{r}'], st['sol'][3], t['du']);    _apply(ws[f'E{r}'], st['sol'][4], t['sh'])
    util_cell = f'C{r+1}'; first = True
    for t in tramos:
        r += 1; ws.row_dimensions[r].height = 25.25
        _apply(ws[f'A{r}'], st['uti'][0], 'Utilizadas' if first else None)
        _apply(ws[f'B{r}'], st['uti'][1], t['serie'])
        _apply(ws[f'C{r}'], st['uti'][2], util_total if first else None); first = False
        _apply(ws[f'D{r}'], st['uti'][3], t['du']); _apply(ws[f'E{r}'], st['uti'][4], t['hu'])
    r += 1; rot_row = r; ws.row_dimensions[r].height = 25.25
    _apply(ws[f'A{r}'], st['rot'][0], 'Rotas'); _apply(ws[f'B{r}'], st['rot'][1], tramos[0]['serie'])
    _apply(ws[f'C{r}'], st['rot'][2]); _apply(ws[f'D{r}'], st['rot'][3]); _apply(ws[f'E{r}'], st['rot'][4])
    first = True
    for t in tramos:
        r += 1; dev_rows.append(r); ws.row_dimensions[r].height = 25.25
        _apply(ws[f'A{r}'], st['dev'][0], 'Devueltas')
        _apply(ws[f'B{r}'], st['dev'][1], t['serie']); _apply(ws[f'C{r}'], st['dev'][2], f'=+E{r}-D{r}')
        _apply(ws[f'D{r}'], st['dev'][3], t['hu']);   _apply(ws[f'E{r}'], st['dev'][4], t['sh'])
    solC = '+'.join(f'C{x}' for x in sol_rows); devC = '+'.join(f'C{x}' for x in dev_rows)
    ws[f'C{rot_row}'] = f'=({solC})-{util_cell}-({devC})'
    r += 1; ws.row_dimensions[r].height = 25.25
    _apply(ws[f'A{r}'], st['tot'][0], 'Total'); _apply(ws[f'B{r}'], st['tot'][1], tramos[0]['serie'])
    _apply(ws[f'C{r}'], st['tot'][2], f'=({solC})'); _apply(ws[f'D{r}'], st['tot'][3]); _apply(ws[f'E{r}'], st['tot'][4])
    r += 2; fr = r
    ws.merge_cells(start_row=fr, start_column=1, end_row=fr+2, end_column=10)
    _apply(ws[f'A{fr}'], foot_s,
           '* FAJAS CORRESPONDIENTES A LAS SERIES Y ANEXOS:\n SERIE _______  Nº: _____________________________')
    for rr in range(fr, fr+3): ws.row_dimensions[rr].height = 24
    r = fr + 5
    ws[f'B{r}'] = '      …………………………………………………………'; ws[f'H{r}'] = '…………………………………………………………'
    _apply(ws[f'A{r+1}'], sig[0], 'Autorización Servicio de Inspección')
    _apply(ws[f'A{r+2}'], sig[1], 'Veterinaria - Firma y sello aclaratorio')
    ws.merge_cells(start_row=r+1, start_column=1, end_row=r+1, end_column=3)
    ws.merge_cells(start_row=r+2, start_column=1, end_row=r+2, end_column=3)
    _apply(ws[f'H{r+1}'], sig[3], 'Firma y sello aclaratorio')
    _apply(ws[f'H{r+2}'], sig[4], 'Responsable de la empresa')
    last = r + 2
    ws.print_area = f'A1:J{last}'
    ws.page_setup.orientation = 'portrait'; ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.title = f'{dd}-{mm}'
    wb.save(outpath)

# --------------------------------------------------------------------------- #
def collect_files(inputs):
    files = []
    for p in inputs:
        if os.path.isdir(p):
            files += sorted(glob.glob(os.path.join(p, '*.xlsx')))
        else:
            files.append(p)
    return files

def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description='Genera Anexos V de Fajas UE por dia.')
    ap.add_argument('inputs', nargs='+', help='Bajada(s) .xlsx o carpeta')
    ap.add_argument('--plantilla', default=os.path.join(here, 'plantilla_Fajas_UE.xlsx'))
    ap.add_argument('--salida', default=os.path.join(here, 'salida'))
    args = ap.parse_args()

    files = collect_files(args.inputs)
    print(f'Bajadas: {len(files)}')
    df = pd.concat([pd.read_excel(f, sheet_name='ag-grid', dtype=str) for f in files],
                   ignore_index=True)
    xlsx_dir = os.path.join(args.salida, 'xlsx'); os.makedirs(xlsx_dir, exist_ok=True)
    dates = sorted({d for d in df['Fecha'].dropna().unique() if isinstance(d, str)})
    hechos = []
    for fecha in dates:
        d = df[df['Fecha'] == fecha]
        tramos, util = day_tramos(d)
        if not tramos:
            continue
        yy, mm, dd = fecha.split('-')
        out = os.path.join(xlsx_dir, f'Anexo_V_Fajas_UE_{dd}-{mm}-{yy}.xlsx')
        build_day(args.plantilla, fecha, tramos, util, out)
        hechos.append((fecha, len(tramos), util, out))
        print(f'  {fecha}: {len(tramos)} tramo(s), {util} utilizadas')
    print(f'Anexos generados: {len(hechos)}')

    # PDF cronologico (requiere LibreOffice)
    pdf_dir = os.path.join(args.salida, '_pdf'); os.makedirs(pdf_dir, exist_ok=True)
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if soffice:
        subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'pdf',
                        '--outdir', pdf_dir] + [h[3] for h in hechos],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            import pypdf
            def key(p):
                dd, mm, yy = os.path.basename(p).replace('.pdf','').split('_')[-1].split('-')
                return (int(yy), int(mm), int(dd))
            pdfs = sorted(glob.glob(os.path.join(pdf_dir, '*.pdf')), key=key)
            w = pypdf.PdfWriter()
            for p in pdfs:
                for pg in pypdf.PdfReader(p).pages: w.add_page(pg)
            w.write(os.path.join(args.salida, 'Anexos_V_Fajas_UE_cronologico.pdf'))
            print(f'PDF cronologico: {len(pdfs)} paginas')
        except Exception as e:
            print('No se pudo unir el PDF:', e)
    else:
        print('AVISO: LibreOffice no encontrado, se omite el PDF.')

    # ZIP por mes con nombres ISO
    staging = os.path.join(args.salida, '_staging'); shutil.rmtree(staging, ignore_errors=True)
    for f in glob.glob(os.path.join(xlsx_dir, '*.xlsx')):
        dd, mm, yy = os.path.basename(f).replace('.xlsx','').split('_')[-1].split('-')
        folder = os.path.join(staging, MESES.get(mm, mm)); os.makedirs(folder, exist_ok=True)
        shutil.copy(f, os.path.join(folder, f'Anexo_V_Fajas_UE_{yy}-{mm}-{dd}.xlsx'))
    shutil.make_archive(os.path.join(args.salida, 'Anexos_V_Fajas_UE_editables'),
                        'zip', staging)
    shutil.rmtree(pdf_dir, ignore_errors=True); shutil.rmtree(staging, ignore_errors=True)
    print('Listo ->', args.salida)

if __name__ == '__main__':
    main()
